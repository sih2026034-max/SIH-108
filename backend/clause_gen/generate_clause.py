import os
import json
import logging
from docx import Document
from typing import List, Dict, Any

logging.basicConfig(level=logging.INFO, format='%(message)s')
logger = logging.getLogger(__name__)

DATASET_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "data", "processed", "normalized_standards.json")

def load_standard_data(standard_ids: List[str]) -> List[Dict[str, Any]]:
    if not os.path.exists(DATASET_PATH):
        logger.warning(f"Dataset not found at {DATASET_PATH}")
        return [{"is_number": sid, "title": "Unknown (Dataset Not Found)"} for sid in standard_ids]
    
    with open(DATASET_PATH, 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    std_map = {item['is_number']: item for item in data}
    results = []
    for sid in standard_ids:
        if sid in std_map:
            results.append(std_map[sid])
        else:
            results.append({"is_number": sid, "title": "Standard title not found in dataset."})
    return results

def format_tender_clause_text(standards_data: List[Dict[str, Any]]) -> str:
    """
    Safely constructs the tender clause string based ONLY on provided facts.
    """
    if not standards_data:
        return "No standard selected."

    clause = "TENDER SPECIFICATION CLAUSE\n\n1. APPLICABLE STANDARDS\nThe items supplied under this contract shall conform to the following Indian Standards:\n"
    
    all_refs = set()
    for i, std in enumerate(standards_data):
        title = std.get('title', 'Specification')
        clause += f"   1.{i+1} {std['is_number']} ({title})\n"
        
        # Collect normative references if they exist
        cross_refs = std.get('cross_references', [])
        if isinstance(cross_refs, list):
            for ref in cross_refs:
                all_refs.add(ref)
        
    clause += "\n2. NORMATIVE REFERENCES\nThe following referenced standards are indispensable for the application of this document:\n"
    
    if all_refs:
        for ref in sorted(all_refs):
            clause += f"   - {ref}\n"
    else:
        clause += "   - (No specific normative references identified in dataset. Refer to section 1 standards.)\n"
        
    clause += "\n3. CERTIFICATION REQUIREMENTS\n"
    # Check if any selected standard requires certification based on data
    cert_required = False
    for std in standards_data:
        req = std.get('certification_requirements')
        if req:
            cert_required = True
            
    if cert_required:
        clause += "   All items shall bear the ISI Mark as per the relevant Quality Control Order and BIS requirements.\n"
        clause += "   Bidders shall furnish a valid BIS licence certificate.\n\n"
    else:
        clause += "   Verify specific BIS certification requirements as per the latest Quality Control Orders (QCO) issued by the relevant Ministry.\n\n"
    
    clause += "4. ACCEPTANCE CRITERIA\n"
    clause += "   [To be completed by the procuring officer based on specific project requirements]\n"
    
    return clause

def generate_docx(clause_text: str, output_path: str):
    """
    Generates a properly formatted Word Document for the tender clause.
    """
    doc = Document()
    doc.add_heading('Technical Specification Clause', 0)
    
    for section in clause_text.split('\n\n'):
        if section.startswith('TENDER'): continue
        if section.strip():
            # Bold the headers
            if section[0].isdigit() and ' ' in section:
                lines = section.split('\n')
                p = doc.add_paragraph()
                p.add_run(lines[0]).bold = True
                if len(lines) > 1:
                    p.add_run('\n' + '\n'.join(lines[1:]))
            else:
                doc.add_paragraph(section)
                
    doc.save(output_path)
    logger.info(f"DOCX saved to {output_path}")
    return output_path

def create_clause(standard_ids: List[str]):
    """
    Main interface for Agent D (API).
    """
    # 1. Fetch metadata for these IDs from dataset
    real_data = load_standard_data(standard_ids)
    
    # 2. Format the text
    text_preview = format_tender_clause_text(real_data)
    
    # 3. Generate DOCX
    os.makedirs(os.path.join(os.path.dirname(__file__), "examples"), exist_ok=True)
    
    # Handle multiple standard IDs in filename safely
    safe_name = str(standard_ids[0]).replace(':', '_').replace('/', '_').replace(' ', '_') if standard_ids else "generic"
    docx_filename = f"clause_{safe_name}.docx"
    docx_path = os.path.join(os.path.dirname(__file__), "examples", docx_filename)
    generate_docx(text_preview, docx_path)
    
    return {
        "text_preview": text_preview,
        "docx_download_url": f"/downloads/{docx_filename}",
        "referenced_is_numbers": [d['is_number'] for d in real_data]
    }

if __name__ == "__main__":
    print(create_clause(["IS 16107 (Part 2/Sec 2):2017"]))
